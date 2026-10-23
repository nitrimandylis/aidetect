"""Builds the demo world for tapes/aidetect.tape: a fake HOME and a short
synthetic Extended Essay draft. Nothing here is real coursework.

    uv run --project . python tapes/fixture.py

Output goes to tapes/fixtures/ (gitignored).
"""

import os

from docx import Document

HERE = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.join(HERE, "fixtures")
HOME = os.path.join(OUT, "home")

SECTIONS = [
    ("Introduction", 1, [
        "I started this essay after a week of watching tide pools at the end of "
        "our street. The same crab sat under the same rock every morning, and "
        "I wanted to know why it never moved. This essay asks how water "
        "temperature changes the feeding rate of the shore crab in a small "
        "rock pool.",
        "My own measurements are limited to one pool and three weeks, so I "
        "treat the results as a first look, not a rule.",
    ]),
    ("Method", 1, [
        "Each morning at low tide I counted how many mussels one crab opened "
        "in thirty minutes. I took the water temperature with a cheap "
        "thermometer and wrote it down before the first count. The crab was "
        "returned to the pool after every trial.",
    ]),
    ("Analysis", 1, [
        "The crab opened more mussels on the warmer mornings. At 11 degrees it "
        "opened two, and at 17 degrees it opened six. The pattern is not "
        "perfect, since a windy day gave me a low count at 16 degrees.",
        "Wind may have stirred up sand and made the crab cautious. I cannot "
        "prove that, because I did not record wind speed.",
    ]),
    ("Conclusion", 1, [
        "Warmer water went with faster feeding in this one pool. A fair next "
        "step is to repeat the counts with several crabs and to log the wind.",
    ]),
    ("Bibliography", 1, [
        "Made-up Author. A Made-up Guide to the Shore. Example Press, 2020.",
    ]),
]


def build_draft(path):
    doc = Document()
    doc.add_paragraph("Extended Essay: Crab Feeding and Water Temperature")
    for title, level, paragraphs in SECTIONS:
        doc.add_heading(title, level=level)
        for text in paragraphs:
            doc.add_paragraph(text)
    doc.save(path)


os.makedirs(HOME, exist_ok=True)
build_draft(os.path.join(HOME, "draft.docx"))
print("fixture written to", OUT)
